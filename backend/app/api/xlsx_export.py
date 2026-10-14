"""Excel (.xlsx) export built from the SAME report dict used by the HTML/PDF/JSON/CSV exports (no recalculation)."""
import io
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SHEETS = ["Executive Summary", "All Findings", "Functional Tests", "Security", "Performance", "Accessibility and UX", "Pages and Coverage", "AI Model Details"]
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")
MAX_CELL = 30000  # Excel hard limit is 32,767 characters
NOT_RECORDED = "Not recorded"


def safe(value):
    """Neutralise formula injection and strip characters Excel cannot store. Numbers/None pass through."""
    if value is None or isinstance(value, (int, float, datetime, bool)):
        return value
    text = _ILLEGAL.sub("", str(value))
    if text.startswith(_FORMULA_PREFIXES):
        text = "'" + text
    return text[:MAX_CELL]


def when(value):
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).replace(tzinfo=None)
    except (ValueError, TypeError):
        return safe(value)


def _evidence(f: dict) -> str:
    return "\n---\n".join(e.get("detail", "") for e in f.get("evidence") or [] if e.get("detail"))


def _step(report: dict, step_id: str) -> dict | None:
    return next((s for s in report.get("steps", []) if s.get("id") == step_id), None)


def _status(step: dict | None) -> str:
    return "Not Run" if not step else {"completed": "Completed", "partial": "Partial", "failed": "Failed", "skipped": "Skipped", "blocked": "Blocked"}.get(step["status"], step["status"].title())


def build_workbook(report: dict, model_context: dict | None = None) -> Workbook:
    wb = Workbook()
    findings = [f for f in report.get("findings", [])]
    steps = report.get("steps", [])
    cases = report.get("test_cases", [])

    def sheet(name, headers, rows, widths, first=False, note=None):
        ws = wb.active if first else wb.create_sheet()
        ws.title = name
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="0B3B36")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        if not rows:
            rows = [[note or "Not Run"] + [""] * (len(headers) - 1)]
        for r in rows:
            ws.append([safe(v) for v in r])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        return ws

    def paint(ws, col_letter, mapping):
        rng = f"{col_letter}2:{col_letter}{max(ws.max_row, 2)}"
        for text, color in mapping.items():
            ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=[f'"{text}"'], fill=PatternFill("solid", bgColor=color)))

    SEV = {"High": "F8C9CC", "Critical": "F2A0A6", "Medium": "FCE4B0", "Low": "CFE0F7"}
    STAT = {"Verified": "C6EFCE", "Possible": "FFEB9C", "Failed": "FFC7CE", "Skipped": "E0E0E0", "Blocked": "E0E0E0", "Not Run": "E0E0E0", "executed": "C6EFCE", "failed": "FFC7CE", "skipped": "E0E0E0", "Passed": "C6EFCE"}

    # 1. Executive summary
    scores = report.get("scores", {})
    exec_cases = [c for c in cases if c.get("status") == "executed"]
    counts = {k: sum(1 for s in steps if s.get("status") == k) for k in ("completed", "partial", "failed", "skipped", "blocked")}
    priorities = report.get("priorities") or []
    top5 = [(p.get("title") or p.get("action") or str(p)) if isinstance(p, dict) else str(p) for p in priorities[:5]]
    rows = [["Target URL", report.get("url")], ["Scan ID", report.get("id")], ["Scan timestamp", when(report.get("created_at"))], ["Finished", when(report.get("finished_at"))],
            ["Overall score (0-100)", report.get("overall_score") if report.get("overall_score") is not None else "Not scored (insufficient coverage)"],
            ["Coverage (%)", report.get("coverage")], ["Report is partial", "Yes" if report.get("partial") else "No"]]
    for key, s in scores.items():
        rows.append([f"Score: {key}", s.get("score") if s.get("score") is not None else f"Not scored ({s.get('status')})"])
    rows += [["Total findings", len(findings)], ["Verified findings", sum(1 for f in findings if f.get("status") == "Verified")], ["Possible findings", sum(1 for f in findings if f.get("status") == "Possible")],
             ["Form test cases executed", len(exec_cases)], ["Form test cases not executed", len(cases) - len(exec_cases)],
             ["Pipeline steps completed", counts["completed"]], ["Pipeline steps partial", counts["partial"]], ["Pipeline steps failed", counts["failed"]],
             ["Pipeline steps skipped", counts["skipped"]], ["Pipeline steps blocked", counts["blocked"]]]
    rows += [[f"Top recommendation {i}", t] for i, t in enumerate(top5, 1)] or [["Top recommendations", "None recorded"]]
    ws = sheet(SHEETS[0], ["Item", "Value"], rows, [34, 90], first=True)
    for r in range(2, ws.max_row + 1):
        if isinstance(ws.cell(r, 2).value, datetime):
            ws.cell(r, 2).number_format = "yyyy-mm-dd hh:mm:ss"

    # 2. All findings
    rows = [[f.get("id"), f.get("title"), f.get("pillar"), f.get("severity"), "High" if f.get("status") == "Verified" else "Medium", f.get("status"),
             (f.get("location") or {}).get("url"), (f.get("location") or {}).get("selector"), f.get("category"), f.get("explanation"), _evidence(f), f.get("suggested_fix"),
             "; ".join(f.get("sources") or []), when(report.get("finished_at"))] for f in findings]
    ws = sheet(SHEETS[1], ["Finding ID", "Title", "Category", "Severity", "Confidence (derived)", "Verification status", "Affected URL", "Element or selector", "Rule ID", "Description", "Evidence", "Recommended fix", "Source scanner", "Detection timestamp"], rows,
               [12, 38, 14, 11, 14, 14, 36, 28, 26, 50, 60, 50, 24, 20], note="No findings recorded")
    for r in range(2, ws.max_row + 1):
        ws.cell(r, 14).number_format = "yyyy-mm-dd hh:mm:ss"
    paint(ws, "D", SEV); paint(ws, "F", STAT)

    # 3. Functional tests (executed form cases + functionality pipeline steps; proposed-but-unexecuted cases are labelled)
    rows = []
    for i, c in enumerate(cases, 1):
        obs = c.get("observations") or []
        rows.append([f"FT-{i:03}", c.get("url"), c.get("name"), f"Form {c.get('form_id')} loaded in the browser; nothing is submitted", f"Enter '{c.get('kind')}' input into the declared fields (client-side validity check only)",
                     c.get("expected"), "; ".join(f"{o.get('selector', '')}: valid={o.get('valid')} {o.get('message', '')}".strip() for o in obs) or c.get("detail") or "No observation recorded",
                     c.get("status"), NOT_RECORDED, c.get("error") or ""])
    for s in steps:
        if s.get("pillar") == "functionality":
            rows.append([f"STEP-{s['id']}", "Site-wide", s.get("title"), "Crawl completed", "Automated check", "Completes without error", s.get("detail"), _status(s), NOT_RECORDED, "" if s.get("status") == "completed" else s.get("detail")])
    ws = sheet(SHEETS[2], ["Test ID", "Page URL", "Test name", "Preconditions", "Action", "Expected result", "Actual result", "Status", "Duration", "Error details"], rows, [12, 34, 30, 34, 40, 20, 50, 12, 14, 30], note="Not Run")
    paint(ws, "H", STAT)

    # 4. Security
    sec = [f for f in findings if f.get("pillar") == "security"]
    auth = "Authorized origin (intrusive checks permitted)" if report.get("authorized") else "Not authorized: read-only checks only"
    rows = [[f.get("id"), (f.get("location") or {}).get("url"), f.get("title"), f.get("severity"), _evidence(f), f.get("status"), f.get("suggested_fix"), "; ".join(f.get("sources") or []), auth] for f in sec]
    for sid in ("custom_security", "zap", "nuclei"):
        st = _step(report, sid)
        if not st or st.get("status") != "completed":
            rows.append([f"STEP-{sid}", "Site-wide", f"Scanner step: {sid}", "", st.get("detail") if st else "Not run", _status(st), "", sid, auth])
    ws = sheet(SHEETS[3], ["Check ID", "Endpoint", "Check name", "Severity", "Evidence", "Verification status", "Remediation", "Scanner", "Authorization or execution limitation"], rows, [14, 36, 36, 11, 60, 16, 50, 22, 40], note="Not Run")
    paint(ws, "D", SEV); paint(ws, "F", STAT)

    # 5. Performance
    http_m = (_step(report, "http_timing") or {}).get("metrics") or {}
    lh = (_step(report, "lighthouse") or {}).get("metrics") or {}
    load = (_step(report, "load") or {}).get("metrics") or {}
    rows = []
    if lh:
        rows.append([report.get("url"), "Lighthouse default profile", lh.get("fcp_ms"), lh.get("lcp_ms"), lh.get("cls"), lh.get("tbt_ms"), lh.get("ttfb_ms"), None, None, "", f"Lighthouse {lh.get('version')}; performance score {lh.get('lighthouse_score')}", lh.get("inp_note") or "Single lab run; not field data"])
    for p in http_m.get("pages", []):
        samples = p.get("samples_ms") or []
        rows.append([p.get("url"), "HTTP timing (no browser)", None, None, None, None, min(samples) if samples else None, p.get("requests"), (p.get("html_bytes") or 0) + (p.get("resource_bytes") or 0), "", f"{len(samples)} runs; fastest shown as TTFB proxy (complete GET time)", "Includes safety-proxy overhead"])
    for r in http_m.get("largest_resources", []):
        rows.append([r.get("url"), "Resource audit", None, None, None, None, r.get("ms"), 1, r.get("bytes"), "Large resource" if r.get("bytes", 0) >= 250_000 else "", "Single GET", f"Referenced from {r.get('page')}"])
    if load:
        rows.append([report.get("url"), "Capped load test", None, None, None, None, load.get("p50_ms"), load.get("requests"), None, "", f"{load.get('users')} users, {load.get('seconds')} s, p95 {load.get('p95_ms')} ms, failures {load.get('failures')}", "Authorized target only; strict caps"])
    ws = sheet(SHEETS[4], ["Page URL", "Device profile", "FCP (ms)", "LCP (ms)", "CLS", "TBT (ms)", "TTFB (ms)", "Request count", "Transferred bytes", "Image findings", "Run count and aggregate measurements", "Measurement notes"], rows, [36, 26, 11, 11, 9, 11, 11, 13, 16, 16, 44, 40],
               note="Not Run" if not (lh or http_m or load) else "No measurements")
    for r in range(2, ws.max_row + 1):
        for col in (3, 4, 6, 7):
            ws.cell(r, col).number_format = "0.0"
        ws.cell(r, 5).number_format = "0.000"; ws.cell(r, 9).number_format = "#,##0"

    # 6. Accessibility and UX
    acc = [f for f in findings if f.get("pillar") == "accessibility"]
    rows = [[(f.get("location") or {}).get("url"), f.get("category"), f.get("title"), f.get("severity"), (f.get("location") or {}).get("selector"), _evidence(f),
             "See rule documentation" if f.get("category") else "Not recorded", f.get("status"), f.get("suggested_fix")] for f in acc]
    ws = sheet(SHEETS[5], ["Page URL", "Rule ID", "Rule name", "Severity", "Affected element", "Evidence", "WCAG reference", "Verification status", "Recommended fix"], rows, [36, 24, 40, 11, 30, 60, 22, 16, 50], note="No accessibility findings recorded" if _step(report, "axe") else "Not Run")
    paint(ws, "D", SEV); paint(ws, "H", STAT)

    # 7. Pages and coverage
    pillar_state = {k: _status(_step(report, sid)) for k, sid in (("sec", "custom_security"), ("perf", "lighthouse"), ("acc", "axe"), ("func", "links"))}
    rows = []
    for p in report.get("pages", []):
        shots = p.get("screenshots") or {}
        rows.append([p.get("url"), f"HTTP {p.get('status_code')}" if p.get("status_code") else "Unknown", "Captured" if shots else "Not captured",
                     f"Site-wide: {pillar_state['func']}", f"Site-wide: {pillar_state['sec']}", f"Site-wide: {pillar_state['perf']}", f"Site-wide: {pillar_state['acc']}", NOT_RECORDED,
                     "" if shots else "No screenshot recorded"])
    ws = sheet(SHEETS[6], ["Discovered URL", "Page status", "Screenshot status", "Functional audit status", "Security status", "Performance status", "Accessibility status", "Retry count", "Failure or skip reason"], rows, [40, 14, 16, 22, 22, 22, 22, 12, 36], note="No pages recorded")

    # 8. AI model details
    ctx = model_context or {}
    catalog = {m.get("ollama_tag"): m for m in ctx.get("catalog", []) if m.get("ollama_tag")}
    used = {}
    for t in report.get("ai_tasks", []):
        used.setdefault(t.get("model"), []).append(t)
    ai_findings = sum(1 for f in findings if f.get("ai_assisted"))
    rows = []
    for model, tasks in used.items():
        meta = catalog.get(model) or {}
        done = sum(1 for t in tasks if t.get("status") == "completed")
        test = (ctx.get("tests") or {}).get(meta.get("id")) or {}
        rows.append([model, meta.get("provider", "Unknown"), f"{model} (exact digest not recorded)", "; ".join(sorted({t.get("name", "").split(" · ")[0] for t in tasks})), (meta.get("license") or {}).get("name", "Unknown \u2014 verify before production"),
                     meta.get("cost_type", "Unknown"), meta.get("deployment", "Unknown"), f"{done}/{len(tasks)} tasks completed" + ("" if done else "; " + (tasks[0].get("detail") or "")),
                     ("Test passed" if test.get("ok") else "Test failed: " + test.get("detail", "")) if test else "Not tested", ai_findings, "; ".join(meta.get("limitations") or ["Unknown"])])
    ws = sheet(SHEETS[7], ["Selected model", "Provider", "Exact version", "Task", "License", "Cost type", "Deployment type", "Availability", "Health-check status", "AI recommendations generated", "Model limitations"], rows, [22, 20, 30, 40, 14, 14, 14, 36, 26, 14, 50],
               note="Not Run (no AI steps recorded)")
    return wb


def render_xlsx(report: dict, model_context: dict | None = None) -> bytes:
    buffer = io.BytesIO()
    build_workbook(report, model_context).save(buffer)
    return buffer.getvalue()
