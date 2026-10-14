#!/usr/bin/env python3
"""Real-browser UI check for the AI Models page and Excel export (needs data/demo-run from run_demo.py).
Run: python scripts/check_ui_models.py   -> prints PASS/FAIL per check, exit 1 on any failure; saves screenshots/actual/."""
import io, json, os, sys, tempfile, time, urllib.request
from pathlib import Path
from playwright.sync_api import sync_playwright
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(Path(__file__).parent))
from _proc import popen_background, stop_pid  # noqa: E402

PORT, OUT = 8092, ROOT / "screenshots" / "actual"
OUT.mkdir(parents=True, exist_ok=True)
data_dir = Path(tempfile.mkdtemp())
src = ROOT / "data" / "demo-run"
if src.exists():
    import shutil
    shutil.copytree(src, data_dir, dirs_exist_ok=True)
env = {**os.environ, "UWT_DATA_DIR": str(data_dir), "UWT_AI_ENABLED": "true", "UWT_OLLAMA_URL": "http://127.0.0.1:1", "UWT_ALLOWLIST": "http://127.0.0.1:8765"}
proc = popen_background([sys.executable, "-m", "uvicorn", "app.main:app", "--port", str(PORT)], cwd=ROOT / "backend", env=env)
results = []
def check(name, ok, detail=""):
    results.append(ok); print(("PASS " if ok else "FAIL ") + name + (f" - {detail}" if detail else ""), flush=True)
try:
    for _ in range(60):
        try:
            jobs = json.load(urllib.request.urlopen(f"http://127.0.0.1:{PORT}/api/jobs", timeout=2))["jobs"]; break
        except OSError: time.sleep(0.5)
    done = [j for j in jobs if j["status"] == "completed"]
    base = f"http://127.0.0.1:{PORT}"
    with sync_playwright() as pw:
        b = pw.chromium.launch(); ctx = b.new_context(viewport={"width": 1366, "height": 900}, accept_downloads=True); pg = ctx.new_page()
        errors = []; pg.on("pageerror", lambda e: errors.append(str(e))); pg.on("console", lambda m: m.type == "error" and errors.append(m.text))
        pg.goto(base); pg.get_by_role("link", name="AI Models").click(); pg.wait_for_url("**/models")
        pg.get_by_role("heading", name="AI Models").wait_for(); check("nav: AI Models page opens", True)
        pg.screenshot(path=str(OUT / "models-page.png"))
        check("honest Ollama status shown", pg.get_by_text("Local Ollama is not reachable.").is_visible())
        pg.get_by_role("button", name="Rank models").click(); pg.get_by_text("Ranking updated").wait_for()
        check("Rank models works", pg.get_by_text("Not ranked").count() == 0)
        pg.get_by_label("2. Type").select_option("paid"); pg.get_by_text("Hosted OpenAI-compatible").wait_for()
        check("paid filter shows hosted entry, buttons disabled", pg.get_by_role("button", name="Test Model").first.is_disabled())
        pg.get_by_label("2. Type").select_option("open")
        pg.get_by_role("button", name="Select").first.click(); pg.get_by_text("only confirmed working after Test Model passes").wait_for()
        check("Select saves and does not claim it works", True)
        pg.get_by_role("button", name="Test Model").first.click(); pg.get_by_role("status").filter(has_text="reach").wait_for(timeout=15000)
        check("Test Model fails honestly when Ollama is down", "success" not in pg.get_by_role("status").inner_text().lower())
        pg.get_by_role("button", name="Details").first.click(); check("Details expands licence info", pg.get_by_text("Official repository").is_visible())
        pg.screenshot(path=str(OUT / "models-page-after.png"), full_page=True)
        if done:
            pg.goto(f"{base}/reports/{done[0]['id']}"); pg.wait_for_selector("text=Export report"); pg.get_by_text("Export report").click()
            with pg.expect_download() as dl: pg.get_by_role("link", name="Excel (XLSX) report").click()
            path = data_dir / dl.value.suggested_filename; dl.value.save_as(path); wb = load_workbook(path)
            check("XLSX downloads via real button and opens", len(wb.sheetnames) == 8, str(wb.sheetnames))
            for fmt, label in (("pdf", "PDF report"), ("html", "HTML report"), ("json", "JSON report"), ("csv", "CSV report")):
                if not pg.locator("details.export-menu[open]").count(): pg.get_by_text("Export report").click()
                with pg.expect_download() as d2: pg.get_by_role("link", name=label).click()
                p2 = data_dir / d2.value.suggested_filename; d2.value.save_as(p2); check(f"{fmt.upper()} download still works", p2.stat().st_size > 100 and p2.suffix == '.' + fmt)
            pg.screenshot(path=str(OUT / "report-with-xlsx-export.png"))
        else:
            check("report page (needs run_demo data)", False, "no completed job found")
        # existing nav/pages still work
        for name, url in (("History", "/history"), ("Reports", "/reports"), ("New test", "/")):
            pg.goto(base + url); check(f"existing page loads: {name}", pg.locator("main, .page, .workspace").first.is_visible())
        errors = [e for e in errors if "503" not in e]  # expected: Test Model returns 503 while Ollama is unreachable
        check("no unexpected browser console/page errors", not errors, "; ".join(errors)[:200])
        b.close()
finally:
    stop_pid(proc.pid)
print(f"\n{sum(results)}/{len(results)} UI checks passed"); sys.exit(0 if all(results) else 1)
